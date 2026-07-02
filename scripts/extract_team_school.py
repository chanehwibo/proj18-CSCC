# -*- coding: utf-8 -*-
"""v2 extractor: gazetteer-based school detection + team/project name.

Improvements over v1:
  * mask contest-name / rules boilerplate ("全国大学生...大赛", "普通高等学校"...)
    so they stop polluting school detection;
  * gazetteer (historical schools from collected-data.xlsx + curated national
    universities + clean 2026 finds) matched as substrings -> high precision;
  * per-source priority (PDF cover > explicit label > README > docs);
  * clean markdown out of explicit labels;
  * keep raw regex candidates (post-filter) for manual review of schools that
    are not yet in the gazetteer.
Outputs data/inputs/2026-extract-report.json (overwrites v1 report).
"""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "data" / "samples" / "2026"
INPUTS = ROOT / "data" / "inputs"
OUT = INPUTS / "2026-extract-report.json"
XLSX = ROOT.parent / "collected-data.xlsx"

# --------------------------------------------------------------------------
# Gazetteer
# --------------------------------------------------------------------------
CURATED = """
清华大学 北京大学 北京航空航天大学 北京理工大学 北京科技大学 北京交通大学 北京邮电大学
北京工业大学 北京化工大学 北京师范大学 北京工商大学 中国人民大学 中国农业大学
中国科学院大学 中国科学技术大学 中国石油大学 中国地质大学 中国矿业大学 中国传媒大学
南开大学 天津大学 天津理工大学 天津科技大学 天津工业大学 河北大学 燕山大学 河北工业大学
太原理工大学 山西大学 内蒙古大学 内蒙古工业大学
大连理工大学 东北大学 辽宁大学 大连海事大学 沈阳工业大学
吉林大学 长春理工大学 东北师范大学 哈尔滨工业大学 哈尔滨工程大学 哈尔滨理工大学 东北林业大学
复旦大学 上海交通大学 同济大学 华东师范大学 上海大学 华东理工大学 上海理工大学 东华大学
上海海事大学 上海海洋大学 上海电力大学 上海科技大学
南京大学 东南大学 南京理工大学 南京航空航天大学 河海大学 江南大学 苏州大学 南京邮电大学
南京信息工程大学 江苏大学 扬州大学 中国矿业大学
浙江大学 浙江工业大学 杭州电子科技大学 杭州师范大学 宁波大学 宁波工程学院 温州大学
中国科学技术大学 合肥工业大学 安徽大学 安徽工业大学
厦门大学 福州大学 华侨大学 福建师范大学 集美大学
南昌大学 江西财经大学 江西理工大学 华东交通大学
山东大学 中国海洋大学 中国石油大学 山东科技大学 青岛大学 齐鲁工业大学 山东师范大学
郑州大学 河南理工大学 河南科技大学 河南大学 郑州轻工业大学 华北水利水电大学
武汉大学 华中科技大学 武汉理工大学 华中师范大学 中国地质大学 武汉科技大学 中南财经政法大学
湖北大学 湖北民族大学 三峡大学 中南民族大学
湖南大学 中南大学 湖南师范大学 湘潭大学 长沙理工大学 湖南工商大学 湖南农业大学 国防科技大学
中山大学 华南理工大学 暨南大学 华南师范大学 深圳大学 广东工业大学 南方科技大学 汕头大学
广州大学 华南农业大学 南方医科大学
广西大学 桂林电子科技大学 广西师范大学
海南大学 重庆大学 重庆邮电大学 重庆理工大学 重庆师范大学 西南大学 西南交通大学 西南财经大学
四川大学 电子科技大学 西南石油大学 成都理工大学 四川师范大学
贵州大学 云南大学 昆明理工大学 云南师范大学
西北工业大学 西安交通大学 西安电子科技大学 西北大学 长安大学 西安理工大学 西安邮电大学
西安建筑科技大学 陕西师范大学 西安科技大学
兰州大学 兰州理工大学 西北师范大学 宁夏大学 青海大学 新疆大学 石河子大学
战略支援部队信息工程大学 信息工程大学 解放军信息工程大学 陆军工程大学 海军工程大学 空军工程大学
"""

# schools whose name is foreign / not a participating unit -> never auto-pick
FOREIGN = {"加州大学", "华盛顿大学", "乔治华盛顿大学", "美国乔治华盛顿大学",
           "麻省理工学院", "斯坦福大学", "剑桥大学", "牛津大学", "香港大学",
           "香港科技大学", "台湾大学"}

# tokens ending in a school-suffix that are NOT schools (boilerplate / rules)
BLACKLIST_SUBSTR = ("全国大学", "普通高等学校", "全国高等学校", "高等学校", "高校大学",
                    "所在学校", "同一学校", "不同学校", "原参赛学校", "参赛学校",
                    "队伍名称和学校", "上游仓库中的学校")
BLACKLIST_EXACT = {"学院", "学校", "大学", "本学院", "各学院", "计算机学院", "软件学院",
                   "信息学院", "网络空间安全学院", "人工智能学院", "电子信息学院",
                   "计算机科学与技术学院", "网络工程学院", "国家示范性软件学院",
                   "示范性软件学院", "未来技术学院", "研究生院", "本科生院", "高校",
                   "高等学校", "普通高等学校"}

# contest / rules phrases masked out before school regex
MASK_PHRASES = [
    "全国大学生计算机系统能力大赛", "全国大学生计算机系统能力培养大赛",
    "大学生计算机系统能力大赛", "计算机系统能力大赛", "操作系统设计赛",
    "操作系统设计大赛", "全国大学生", "普通高等学校", "全国高等学校",
    "全国普通高等学校", "参赛对象为", "所在学校", "同一学校", "不同学校",
]

# prefix connective/verb chars to peel off a glued school token
PREFIX_NOISE = re.compile(
    r"^(?:感谢|归属院校|归属|来自|本仓库是|本仓库|本项目|本人完全了解|完全了解|"
    r"队员来自|三位成员以|成员以|学校为|就读于|就读|隶属于|隶属|依托|"
    r"由|为|是|在|的|和|与|及|向|对|让|使|给|本|该|我|们|届|年|全国|面向|参与|"
    r"参加|一个|一款|适配了|项目|系统|当前|正|作为|参赛内核)+"
)

SCHOOL_RE = re.compile(
    r"[一-龥]{2,15}?(?:大学|学院|职业技术学院|职业技术师范大学|师范学院|理工学院)"
    r"(?:[（(][一-龥]{1,6}[)）])?"
)

LABEL_SCHOOL = re.compile(
    r"(?:参赛)?(?:学校|院校|所在学校|所在院校|学校名称|培养单位|参赛单位|单位名称|团队所在学校|高校名称)"
    r"\s*[：:是为]\s*([^\n\r，,。；;、|)）]{2,40})"
)
LABEL_TEAM = re.compile(
    r"(?:队伍名称|队伍名|队名|团队名称|团队名|战队名称|参赛队伍|参赛队名)"
    r"\s*[：:是为]?\s*([^\n\r，,。；;、|)）]{1,40})"
)
LABEL_PROJECT = re.compile(
    r"(?:作品名称|项目名称|作品名|系统名称)\s*[：:是为]\s*([^\n\r，,。；;、|)）]{1,40})"
)
LABEL_MEMBERS = re.compile(
    r"(?:队员|成员|团队成员|小组成员|作者|参赛队员)\s*[：:]\s*([^\n\r]{2,80})"
)



TEAM_BAD_SUBSTR = ("\u8d5b\u9053", "\u6bd4\u8d5b", "\u7ade\u8d5b", "\u64cd\u4f5c\u7cfb\u7edf", "\u5185\u6838", "\u63d0\u4ea4", "\u7f16\u53f7", "\u4fe1\u606f",
                   "\u9875\u9762\u663e\u793a", "\u81f4\u4ee5", "\u8c22\u610f", "Linux", "\u80fd\u591f", "\u6307\u5b9a", "\u57fa\u4e8e", "\u6539\u9020",
                   "\u5168\u56fd\u5927\u5b66\u751f", "\u8bbe\u8ba1", "\u8981\u6c42", "\u4f5c\u54c1", "\u9879\u76ee", "\u5b66\u6821")
SCHOOL_LABELS = ("\u5b66\u6821", "\u9662\u6821", "\u6240\u5728\u5b66\u6821", "\u6240\u5728\u9662\u6821", "\u5b66\u6821\u540d\u79f0",
                 "\u57f9\u517b\u5355\u4f4d", "\u53c2\u8d5b\u5355\u4f4d", "\u5355\u4f4d\u540d\u79f0", "\u56e2\u961f\u6240\u5728\u5b66\u6821", "\u9ad8\u6821\u540d\u79f0")
TEAM_LABELS = ("\u961f\u4f0d\u540d\u79f0", "\u961f\u4f0d\u540d", "\u961f\u540d", "\u56e2\u961f\u540d\u79f0", "\u56e2\u961f\u540d",
               "\u6218\u961f\u540d\u79f0", "\u53c2\u8d5b\u961f\u4f0d", "\u53c2\u8d5b\u961f\u540d")
PROJECT_LABELS = ("\u4f5c\u54c1\u540d\u79f0", "\u9879\u76ee\u540d\u79f0", "\u4f5c\u54c1\u540d", "\u7cfb\u7edf\u540d\u79f0")
MEMBER_LABELS = ("\u961f\u5458", "\u6210\u5458", "\u56e2\u961f\u6210\u5458", "\u5c0f\u7ec4\u6210\u5458", "\u4f5c\u8005", "\u53c2\u8d5b\u961f\u5458")

def load_gazetteer() -> set[str]:
    gaz = set(CURATED.split())
    try:
        import openpyxl
        wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
        for r in list(wb["Sheet1"].iter_rows(values_only=True))[1:]:
            if r[3]:
                gaz.add(str(r[3]).strip())
    except Exception as e:  # noqa
        print("warn: could not read xlsx gazetteer:", e, file=sys.stderr)
    # normalise half/full-width parens variants of 哈工大 depth campuses
    extra = set()
    for s in gaz:
        if "(" in s:
            extra.add(s.replace("(", "（").replace(")", "）"))
        if "（" in s:
            extra.add(s.replace("（", "(").replace("）", ")"))
    return gaz | extra


def clean_label(v: str | None) -> str | None:
    if not v:
        return None
    v = re.sub(r"<[^>]+>", "", v)
    v = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", v)
    v = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", v)
    v = re.sub(r"[*_`#>]", "", v).strip()
    v = re.sub(r"^[-\u2014\uff1a:|\s]+", "", v)
    v = re.split(r"[\uff0c,\u3002\uff1b;\n\r]", v, maxsplit=1)[0]
    v = v.strip(" \t\uff1a:-\u2014|)\uff09(\uff08\u3010\u3011[]")
    return v or None


def extract_markdown_label(text: str, labels: tuple[str, ...]) -> str | None:
    label_alt = "|".join(re.escape(x) for x in labels)
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        plain = re.sub(r"[*_`#>]", "", line)
        m = re.match(rf"^\|?\s*(?:{label_alt})\s*\|\s*([^|]+)", plain)
        if m:
            val = clean_label(m.group(1))
            if val:
                return val
        m = re.match(rf"^[-+*]?\s*(?:{label_alt})\s*[\uff1a:=\u4e3a\u662f]\s*(.+)$", plain)
        if m:
            val = clean_label(m.group(1))
            if val:
                return val
    return None


def valid_team_name(v: str | None) -> str | None:
    v = clean_label(v)
    if not v:
        return None
    if len(v) > 24:
        return None
    if any(b in v for b in TEAM_BAD_SUBSTR):
        return None
    if re.fullmatch(r"T\d{6,}[-\d]*", v, re.I):
        return None
    return v

def normalize_school(s: str | None) -> str | None:
    """Canonicalise a school name: strip glue prefixes, unify campus parens,
    reduce '大学…学院' department suffix to the university, map 国防科技大学."""
    s = clean_label(s)
    if not s:
        return None
    s = s.replace("(", "（").replace(")", "）")
    s = strip_prefix(s)
    if "国防科" in s and ("大学" in s or "学" in s):
        return "国防科技大学"
    if "大学" in s:
        # keep from start through '大学' + optional （campus）; drops trailing 学院/系 etc.
        m = re.match(r"^(.*?大学(?:（[^）]{1,8}）)?)", s)
        if m:
            s = m.group(1)
    if not s or s in BLACKLIST_EXACT or len(s) < 4:
        return None
    if any(b in s for b in BLACKLIST_SUBSTR):
        return None
    return s


def mask_boilerplate(text: str) -> str:
    for p in MASK_PHRASES:
        text = text.replace(p, "　" * len(p))
    return text


def strip_prefix(tok: str) -> str:
    prev = None
    while prev != tok:
        prev = tok
        tok = PREFIX_NOISE.sub("", tok)
    return tok


def regex_school_candidates(text: str) -> dict[str, int]:
    text = mask_boilerplate(text)
    freq: dict[str, int] = {}
    for m in SCHOOL_RE.finditer(text):
        tok = strip_prefix(m.group(0).strip())
        if not tok or tok in BLACKLIST_EXACT:
            continue
        if any(b in tok for b in BLACKLIST_SUBSTR):
            continue
        if len(tok) < 4:
            continue
        freq[tok] = freq.get(tok, 0) + 1
    return dict(sorted(freq.items(), key=lambda kv: -kv[1]))


def gazetteer_hits(text: str, gaz: set[str]) -> list[str]:
    text = mask_boilerplate(text)
    hits = []
    for s in gaz:
        if s in FOREIGN:
            continue
        if s in text:
            hits.append(s)
    # prefer longer, more specific names (哈尔滨工业大学（深圳） over 哈尔滨工业大学)
    hits.sort(key=len, reverse=True)
    # dedupe by containment: drop a hit that is a substring of a longer hit
    out = []
    for h in hits:
        if not any(h != o and h in o for o in hits):
            out.append(h)
    return out


# words near a school mention that signal a CITATION of an upstream / other work
CITE_CTX = ("参考", "基于", "感谢", "致谢", "借鉴", "上游", "fork", "Fork", "FORK",
            "rcore", "rCore", "RCore", "starry", "Starry", "STARRY", "arceos",
            "ArceOS", "ARCEOS", "开源训练营", "训练营", "联合开发", "优秀参赛作品",
            "优秀作品", "二次开发", "移植自", "源自", "改编自", "衍生", "xv6", "ByteOS",
            "byteos", "NPUcore", "npucore", "参赛内核参与")
# words near a school mention that signal the TEAM'S OWN affiliation
AFFIL_CTX = ("学校", "院校", "队伍", "队员", "成员", "团队", "我们", "本人", "来自",
             "就读", "隶属", "归属", "单位", "作者", "项目人员", "参赛", "组员",
             "培养单位", "所在")
# files that are contest mirrors / leaderboards / charters -> never own-school evidence
EXCLUDE_FILE_KW = ("rank", "dashboard", "scoring", "snapshot", "leaderboard",
                   "排行", "榜单", "章程", "os-contest-2026", "os-contest2026",
                   "introduction.md", "01-introduction", "赛道说明", "赛题",
                   "getting-started", "评测", "autotest", "官方")


def is_excluded_file(label: str) -> bool:
    low = label.lower()
    return any(kw.lower() in low for kw in EXCLUDE_FILE_KW)


def hit_snippets(text: str, school: str, radius: int = 30) -> list[dict]:
    text = mask_boilerplate(text)
    out = []
    start = 0
    while True:
        idx = text.find(school, start)
        if idx < 0:
            break
        a = max(0, idx - radius)
        b = min(len(text), idx + len(school) + radius)
        win = text[a:b]
        snip = re.sub(r"\s+", " ", win).strip()
        cite = any(c in win for c in CITE_CTX)
        affil = any(c in win for c in AFFIL_CTX)
        out.append({"snippet": snip, "cite": cite, "affil": affil})
        start = idx + len(school)
        if len(out) >= 4:
            break
    return out


def readme_h1(text: str) -> str | None:
    for line in text.splitlines():
        s = line.strip()
        if s.startswith("# "):
            title = s[2:].strip().strip("#").strip()
            title = re.sub(r"!\[[^\]]*\]\([^)]*\)", "", title).strip()
            title = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", title).strip()
            title = re.sub(r"<[^>]+>", "", title).strip()
            if title:
                return title
    return None


def read_text_file(p: Path, limit: int = 80_000) -> str:
    try:
        return p.read_text(encoding="utf-8", errors="ignore")[:limit]
    except OSError:
        return ""


def pdf_first_pages(p: Path, pages: int = 3, limit: int = 8000) -> str:
    try:
        import fitz
        doc = fitz.open(p)
    except Exception:
        return ""
    out = []
    try:
        for i in range(min(pages, doc.page_count)):
            out.append(doc.load_page(i).get_text("text"))
    except Exception:
        pass
    finally:
        try:
            doc.close()
        except Exception:
            pass
    return "\n".join(out)[:limit]


def docx_text(p: Path, limit: int = 8000) -> str:
    try:
        import docx
        d = docx.Document(str(p))
    except Exception:
        return ""
    return "\n".join(par.text for par in d.paragraphs)[:limit]


def logo_hints(readme: str) -> list[str]:
    out = []
    for m in re.finditer(r"!\[([^\]]*)\]\(([^)]+)\)", readme):
        alt, src = m.group(1).strip(), m.group(2).strip()
        name = src.rsplit("/", 1)[-1]
        if alt:
            out.append(alt)
        if any(k in src.lower() for k in ("logo", "school", "univ", "校", "badge")):
            out.append(f"[img]{name}")
    # also raw <img src=...>
    for m in re.finditer(r"<img[^>]+src=[\"']([^\"']+)[\"']", readme):
        src = m.group(1)
        if any(k in src.lower() for k in ("logo", "school", "univ", "校", "badge")):
            out.append(f"[img]{src.rsplit('/',1)[-1]}")
    return out


def gather(repo: Path, gaz: set[str]) -> dict:
    texts: list[tuple[str, str]] = []
    readme = ""
    for name in ("README.md", "README", "README.txt", "readme.md", "Readme.md", "README.MD"):
        rp = repo / name
        if rp.exists():
            readme = read_text_file(rp)
            texts.append(("README", readme))
            break

    md_txt = 0
    for f in repo.rglob("*"):
        if md_txt >= 60:
            break
        if any(part in {".git", "node_modules", "target", "vendor"} for part in f.parts):
            continue
        if f.is_file() and f.suffix.lower() in (".md", ".txt") and f.name.lower() not in ("readme.md",) \
                and f.stat().st_size < 300_000:
            rel = str(f.relative_to(repo))
            if rel == "README.md":
                continue
            texts.append((rel, read_text_file(f, 40_000)))
            md_txt += 1

    pdfs = []
    n_pdf = 0
    pdf_texts: list[tuple[str, str]] = []
    for f in sorted(repo.rglob("*.pdf")):
        if n_pdf >= 10:
            break
        if any(part in {".git"} for part in f.parts):
            continue
        t = pdf_first_pages(f)
        rel = str(f.relative_to(repo))
        pdfs.append(rel)
        n_pdf += 1
        if t.strip():
            pdf_texts.append((rel, t))
    n_docx = 0
    docx_texts: list[tuple[str, str]] = []
    for f in sorted(repo.rglob("*.docx")):
        if n_docx >= 6 or f.name.startswith("~$"):
            continue
        t = docx_text(f)
        if t.strip():
            docx_texts.append((str(f.relative_to(repo)), t))
        n_docx += 1

    # priority sources for school: PDF covers + docx first (formal docs), then README, then other md
    ordered = pdf_texts + docx_texts + [("README", readme)] + \
        [(l, t) for l, t in texts if l != "README"]

    # ---- explicit "label: value" fields (search priority-ordered sources) ----
    all_text = "\n".join(t for _, t in ordered)
    regex_cands = regex_school_candidates(all_text)

    def first_label(rx, src_list, labels: tuple[str, ...] | None = None, *, team: bool = False):
        for label, t in src_list:
            if is_excluded_file(label):
                continue
            val = extract_markdown_label(t, labels) if labels else None
            if not val:
                m = rx.search(t)
                if m:
                    val = clean_label(m.group(1))
            if val:
                val = valid_team_name(val) if team else clean_label(val)
                if val:
                    return val, label
        return None, None

    label_school, label_school_src = first_label(LABEL_SCHOOL, ordered, SCHOOL_LABELS)
    label_team, _ = first_label(LABEL_TEAM, ordered, TEAM_LABELS, team=True)
    label_project, _ = first_label(LABEL_PROJECT, ordered, PROJECT_LABELS)
    label_members, _ = first_label(LABEL_MEMBERS, ordered, MEMBER_LABELS)

    # ---- school detection: score each gazetteer hit by occurrence context ----
    # affiliation mention = +3, neutral (non-cite, non-excluded file) = +1,
    # citation = 0 (ignored). Excluded files (leaderboards/charters) never count.
    school_scores: dict[str, dict] = {}
    for label, t in ordered:
        excluded = is_excluded_file(label)
        for h in gazetteer_hits(t, gaz):
            snips = hit_snippets(t, h)
            rec = school_scores.setdefault(
                h, {"score": 0, "sources": set(), "snippets": [], "affil": False})
            for s in snips:
                if excluded:
                    continue
                if s["affil"] and not s["cite"]:
                    rec["score"] += 3
                    rec["affil"] = True
                elif not s["cite"]:
                    rec["score"] += 1
                if len(rec["snippets"]) < 4:
                    rec["snippets"].append({"src": label, "excluded": excluded, **s})
            rec["sources"].add(label)

    def src_rank(sources) -> int:
        for i, (label, _) in enumerate(ordered):
            if label in sources:
                return i
        return 999

    ranked = sorted(
        ((k, v) for k, v in school_scores.items() if v["score"] > 0),
        key=lambda kv: (-kv[1]["score"], src_rank(kv[1]["sources"])),
    )

    school = None
    school_src = None
    confidence = "none"
    if ranked:
        school = normalize_school(ranked[0][0])
        school_src = sorted(ranked[0][1]["sources"], key=lambda s: src_rank({s}))[0]
        confidence = "high" if ranked[0][1]["affil"] else "medium"

    # explicit self-label is the strongest signal; trust it even off-gazetteer.
    if label_school:
        cand = normalize_school(label_school)
        if cand and (cand.endswith(("大学", "学院")) or gazetteer_hits(cand, gaz)):
            school = cand
            school_src = f"label:{label_school_src}"
            confidence = "high"

    needs_vision = (school is None) and (bool(pdfs) or bool(logo_hints(readme)))

    return {
        "readme_h1": readme_h1(readme),
        "school": school,
        "school_confidence": confidence,
        "school_source": school_src,
        "school_all_hits": {
            k: {"score": v["score"], "affil": v["affil"],
                "sources": sorted(v["sources"], key=lambda s: src_rank({s})),
                "snippets": v["snippets"]}
            for k, v in ranked
        },
        "label_school": label_school,
        "label_team": label_team,
        "label_project": label_project,
        "label_members": label_members,
        "regex_school_candidates": regex_cands,
        "logo_hints": logo_hints(readme),
        "needs_vision": needs_vision,
        "pdfs": pdfs,
        "has_pdf_text": [l for l, _ in pdf_texts],
        "sources_scanned": [l for l, _ in ordered if _.strip()],
    }


def main() -> int:
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    gaz = load_gazetteer()
    print(f"gazetteer size: {len(gaz)}")
    inputs = json.loads((INPUTS / "2026-site-inputs.json").read_text(encoding="utf-8"))
    items = inputs["2026"]
    report = {}
    for i, it in enumerate(items, 1):
        repo_id = it["repo_id"]
        repo = SAMPLES / repo_id
        if not repo.exists():
            report[repo_id] = {"error": "missing repo dir", "entry_no": it.get("entry_no")}
            continue
        info = gather(repo, gaz)
        info["entry_no"] = it.get("entry_no")
        info["url"] = it.get("url")
        report[repo_id] = info
        print(f"[{i}/{len(items)}] {repo_id[:42]:42s} school={str(info['school']):20s} "
              f"h1={info['readme_h1']!r}")
    OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")

    got = sum(1 for v in report.values() if v.get("school"))
    high = sum(1 for v in report.values() if v.get("school_confidence") == "high")
    med = sum(1 for v in report.values() if v.get("school_confidence") == "medium")
    vision = sum(1 for v in report.values() if v.get("needs_vision"))
    print(f"\nwrote {OUT}")
    print(f"school resolved: {got}/{len(report)}  (high={high}, medium={med})")
    print(f"needs vision (no text school but has pdf/logo): {vision}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
